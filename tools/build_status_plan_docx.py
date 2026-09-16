"""Build the 2.0 status and execution plan DOCX from its Markdown source."""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "20_Demo_2.0当前完成情况与后续执行计划.md"
OUTPUT = SOURCE.with_suffix(".docx")
TEAL = "0B6F65"
LIGHT_TEAL = "E8F3F1"
PALE = "F5F8F7"
GRID = "D9E2E0"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=GRID, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text, default_size=10.5, color="263B38"):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=default_size, bold=True, color="111111")
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=max(8.5, default_size - 1), color=TEAL)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=default_size, color=color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    widths = [1 / len(rows[0])] * len(rows[0])
    if len(rows[0]) == 3:
        widths = [0.18, 0.52, 0.30]
    elif len(rows[0]) == 5:
        widths = [0.16, 0.20, 0.24, 0.27, 0.13]
    table_font_size = 8.0 if len(rows[0]) >= 5 else 8.8
    for i, text in enumerate(rows[0]):
        cell = header.cells[i]
        cell.width = Inches(6.25 * widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, TEAL)
        set_cell_border(cell)
        set_cell_margin(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, text, 9.2, "FFFFFF")
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row_data in enumerate(rows[1:], 1):
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cell = cells[i]
            cell.width = Inches(6.25 * widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_fill(cell, "FFFFFF" if row_index % 2 else PALE)
            set_cell_border(cell)
            set_cell_margin(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, text, table_font_size)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_inline(paragraph, "消费者洞察 Demo 2.0   ", 8.5, "6D7E7A")
    run = paragraph.add_run("第 ")
    set_font(run, size=8.5, color="6D7E7A")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, size=8.5, color="6D7E7A")


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("263B38")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.35

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (("Heading 1", 16, 18, 8), ("Heading 2", 12.5, 13, 6)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_inline(header, "项目状态与执行计划   2026 年 9 月 16 日", 8.5, "71837F")
    add_page_number(section.footer.paragraphs[0])


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    in_code = False
    code_lines = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.35)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.keep_together = True
                run = p.add_run("\n".join(code_lines))
                set_font(run, name="Microsoft YaHei", size=8.4, color="2D4D47")
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s|:-]+\|$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                values = [part.strip() for part in lines[i].strip().strip("|").split("|")]
                if i == 0 or not all(re.fullmatch(r"[:\- ]+", value) for value in values):
                    rows.append(values)
                i += 1
            if len(rows) > 1:
                add_table(doc, rows)
            continue
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p_ppr = p._p.get_or_add_pPr()
            p_border = p_ppr.find(qn("w:pBdr"))
            if p_border is not None:
                p_ppr.remove(p_border)
            add_inline(p, line[2:], 25, "000000")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            meta = doc.add_paragraph()
            add_inline(meta, "当前完成情况  验证证据  后续执行步骤", 10, TEAL)
            meta.paragraph_format.space_after = Pt(15)
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:], 16, "000000")
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], 12.5, "000000")
        elif re.match(r"^\d+\. ", line):
            match = re.match(r"^(\d+)\. (.*)$", line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            add_inline(p, f"{match.group(1)}. ", 10.5, "263B38")
            add_inline(p, match.group(2))
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            p.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    props = doc.core_properties
    props.title = "消费者洞察 Demo 2.0 当前完成情况与后续执行计划"
    props.subject = "2.0 已完成工作 验证证据 后续执行步骤"
    props.author = "消费者洞察 Demo 项目组"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
