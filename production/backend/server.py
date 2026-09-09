"""Authenticated personal research API. No uploaded text is trusted as markup."""
import csv
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sqlite3
import tempfile
import time
from contextlib import asynccontextmanager
from types import MappingProxyType
from urllib.parse import urlparse
from pathlib import Path

from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Form
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, Response, FileResponse
from pydantic import BaseModel, Field
try:
    from . import db
except ImportError:
    import db

MAX_ROWS = 5000
_USERNAME = re.compile(r'^[A-Za-z0-9_\-\u3400-\u4dbf\u4e00-\u9fff]+$')

def fail(message, status=400):
    raise HTTPException(status, message)

def normalize_username(username):
    display_name = username.strip()
    if not 3 <= len(display_name) <= 32 or not _USERNAME.fullmatch(display_name):
        raise ValueError('用户名需为3至32位中文、字母、数字、下划线或短横线。')
    return display_name.casefold(), display_name

def hash_password(password):
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode(), salt=salt, n=16384, r=8, p=1)
    return f'scrypt$16384$8$1${salt.hex()}${digest.hex()}'

def verify_password(password, encoded):
    try:
        if encoded.startswith('scrypt$'):
            _, n, r, p, salt, expected = encoded.split('$')
            actual = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), n=int(n), r=int(r), p=int(p)).hex()
        else:
            salt, expected = encoded.split(':')
            actual = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), n=16384, r=8, p=1).hex()
        return hmac.compare_digest(actual, expected)
    except (TypeError, ValueError):
        return False

@asynccontextmanager
async def lifespan(app):
    db.init()
    password = os.environ.get('RADAR_ADMIN_PASSWORD', '')
    with db.connect() as c:
        stored = c.execute("SELECT value FROM meta WHERE key='password'").fetchone()
        if not stored:
            if len(password) < 12:
                raise RuntimeError('首次启动需配置至少12字符 RADAR_ADMIN_PASSWORD。')
            c.execute('INSERT INTO meta VALUES (?,?)', ('password', hash_password(password)))
    yield

app = FastAPI(title='Consumer Signal Radar', lifespan=lifespan)

@app.exception_handler(RequestValidationError)
async def validation_error_without_secrets(request: Request, exc: RequestValidationError):
    return JSONResponse({'detail':'请求参数格式无效。'}, status_code=422)

@app.middleware('http')
async def security(request: Request, call_next):
    if request.url.path.startswith('/api/'):
        if request.method not in ('GET','HEAD','OPTIONS'):
            origin = request.headers.get('origin')
            allowed = {str(request.base_url).rstrip('/'), *filter(None, os.environ.get('RADAR_ALLOWED_ORIGINS','').split(','))}
            if (origin and origin not in allowed) or request.headers.get('sec-fetch-site') == 'cross-site':
                return JSONResponse({'detail':'禁止跨站写入请求。'},403)
        if request.url.path not in ('/api/auth/register','/api/auth/login','/api/auth/admin-login','/api/health'):
            raw = request.cookies.get('radar_session','')
            token = hashlib.sha256(raw.encode()).hexdigest()
            with db.connect() as c:
                session = c.execute('''SELECT s.user_id,s.role,u.display_name,u.status
                    FROM sessions s LEFT JOIN users u ON u.id=s.user_id
                    WHERE s.token=? AND s.expires>?''',(token,time.time())).fetchone()
            if not session or (session['role']=='user' and (not session['user_id'] or session['status']!='active')):
                return JSONResponse({'detail':'请先登录。'},401)
            identity = {
                'id': session['user_id'] if session['role']=='user' else 'admin',
                'username': session['display_name'] if session['role']=='user' else '管理员',
                'role': session['role'],
            }
            request.state.identity = MappingProxyType(identity)
            resource = request.url.path
            if resource == '/api/settings' and request.method == 'POST' and identity['role'] != 'admin':
                return JSONResponse({'detail':'仅管理员可以修改模型配置。'},403)
            project_match = re.match(r'^/api/projects/([^/]+)(?:/|$)', resource)
            job_match = re.match(r'^/api/jobs/([^/]+)(?:/|$)', resource)
            report_match = re.match(r'^/api/reports/([^/]+)(?:/|$)', resource)
            write = request.method not in ('GET','HEAD','OPTIONS')
            try:
                with db.connect() as c:
                    if project_match:
                        copy_action = request.method == 'POST' and resource.endswith('/copy')
                        _visible_project(c,project_match.group(1),identity,write=write and not copy_action)
                    elif job_match:
                        get_job(c,job_match.group(1),identity,write=write)
                    elif report_match:
                        _report_for(c,report_match.group(1),identity)
            except HTTPException as exc:
                return JSONResponse({'detail':exc.detail}, exc.status_code)
    response = await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Cache-Control']='no-store'
    return response

@app.get('/api/health')
def health():
    with db.connect() as c:
        c.execute('SELECT 1')
    return {'status':'ok','service':'consumer-signal-radar'}

class UserCredentials(BaseModel):
    username: str = Field(min_length=1, max_length=64)
    password: str = Field(min_length=8, max_length=128)

class AdminCredentials(BaseModel):
    password: str = Field(min_length=1, max_length=128)

def _identity(user_id, username, role):
    return {'id':user_id, 'username':username, 'role':role}

def _new_session(connection, identity, previous_raw=''):
    now = time.time()
    connection.execute('DELETE FROM sessions WHERE expires<?', (now,))
    if previous_raw:
        connection.execute('DELETE FROM sessions WHERE token=?', (hashlib.sha256(previous_raw.encode()).hexdigest(),))
    if identity['role']=='user':
        connection.execute('DELETE FROM sessions WHERE user_id=?', (identity['id'],))
    else:
        connection.execute("DELETE FROM sessions WHERE role='admin'")
    raw = secrets.token_urlsafe(40)
    user_id = identity['id'] if identity['role']=='user' else None
    connection.execute('INSERT INTO sessions(token,expires,user_id,role) VALUES (?,?,?,?)',
                       (hashlib.sha256(raw.encode()).hexdigest(), now+43200, user_id, identity['role']))
    return raw

def _login_response(identity, raw):
    response = JSONResponse(identity)
    response.set_cookie('radar_session',raw,httponly=True,samesite='strict',
                        secure=os.environ.get('RADAR_COOKIE_SECURE','false').lower()=='true',
                        max_age=43200,path='/')
    return response

def _login_allowed(connection, request, now):
    ip = request.client.host if request.client else 'local'
    connection.execute('DELETE FROM attempts WHERE at<?',(now-900,))
    if connection.execute('SELECT count(*) FROM attempts WHERE ip=?',(ip,)).fetchone()[0]>=10:
        fail('登录尝试过多，请15分钟后重试。',429)
    return ip

def _login_failed(connection, ip, now):
    connection.execute('INSERT INTO attempts VALUES (?,?)',(ip,now))
    connection.commit()
    fail('用户名或密码错误。',401)

@app.post('/api/auth/register')
def register(body: UserCredentials, request: Request):
    try:
        username_norm, display_name = normalize_username(body.username)
    except ValueError as exc:
        fail(str(exc))
    user_id = db.uid()
    identity = _identity(user_id, display_name, 'user')
    try:
        with db.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            c.execute('''INSERT INTO users(id,username_norm,display_name,password_hash,role,status,created_at)
                VALUES (?,?,?,?,?,?,?)''',
                (user_id,username_norm,display_name,hash_password(body.password),'user','active',time.time()))
            raw = _new_session(c, identity, request.cookies.get('radar_session',''))
    except sqlite3.IntegrityError:
        fail('用户名已存在。',409)
    return _login_response(identity, raw)

@app.post('/api/auth/login')
def login(body: UserCredentials, request: Request):
    now = time.time()
    try:
        username_norm, _ = normalize_username(body.username)
    except ValueError:
        username_norm = ''
    with db.connect() as c:
        ip = _login_allowed(c, request, now)
        user = c.execute("SELECT * FROM users WHERE username_norm=? AND role='user' AND status='active'",(username_norm,)).fetchone()
        dummy = c.execute("SELECT value FROM meta WHERE key='password'").fetchone()[0]
        if not verify_password(body.password, user['password_hash'] if user else dummy):
            _login_failed(c, ip, now)
        c.execute('DELETE FROM attempts WHERE ip=?',(ip,))
        identity = _identity(user['id'], user['display_name'], 'user')
        raw = _new_session(c, identity, request.cookies.get('radar_session',''))
    return _login_response(identity, raw)

@app.post('/api/auth/admin-login')
def admin_login(body: AdminCredentials, request: Request):
    now = time.time()
    with db.connect() as c:
        ip = _login_allowed(c, request, now)
        saved = c.execute("SELECT value FROM meta WHERE key='password'").fetchone()[0]
        if not verify_password(body.password, saved):
            _login_failed(c, ip, now)
        c.execute('DELETE FROM attempts WHERE ip=?',(ip,))
        identity = _identity('admin', '管理员', 'admin')
        raw = _new_session(c, identity, request.cookies.get('radar_session',''))
    return _login_response(identity, raw)

@app.get('/api/auth/me')
def me(request: Request):
    return dict(request.state.identity)

@app.post('/api/auth/logout')
def logout(request: Request):
    with db.connect() as c:
        c.execute('DELETE FROM sessions WHERE token=?',(hashlib.sha256(request.cookies.get('radar_session','').encode()).hexdigest(),))
    response=JSONResponse({'ok':True}); response.delete_cookie('radar_session'); return response

@app.get('/api/settings')
def settings(request: Request):
    detail = _settings_detail()
    if request.state.identity['role'] != 'admin':
        return {key:detail[key] for key in ('configured','paid_enabled','max_rows')}
    return detail

def _settings_detail():
    api_key = os.getenv('RADAR_API_KEY','')
    base_url = os.getenv('RADAR_BASE_URL','')
    parsed = urlparse(base_url)
    safe_base_url = base_url if parsed.scheme == 'https' and parsed.hostname and not (parsed.username or parsed.password or parsed.query or parsed.fragment) else ''
    return {'model':os.getenv('RADAR_MODEL',''),'embedding_model':os.getenv('RADAR_EMBEDDING_MODEL',''),'base_url':safe_base_url,'key_configured':bool(api_key),'configured':all(os.getenv(k) for k in ('RADAR_API_KEY','RADAR_BASE_URL','RADAR_MODEL','RADAR_EMBEDDING_MODEL')),'paid_enabled':os.getenv('RADAR_PAID_ENABLED','false').lower()=='true','max_rows':MAX_ROWS}

class SettingsUpdate(BaseModel):
    api_key: str = Field(default='', max_length=512)
    base_url: str = Field(min_length=1, max_length=2048)
    model: str = Field(min_length=1, max_length=128)
    embedding_model: str = Field(min_length=1, max_length=128)
    paid_enabled: bool = False

_MODEL_NAME = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._:/-]{0,127}$')
_API_KEY = re.compile(r'^[A-Za-z0-9._~+/=-]{8,512}$')

def _env_path():
    return Path(os.environ.get('RADAR_ENV_PATH', str(Path(__file__).resolve().parent.parent / '.env'))).resolve()

def _replace_env_values(path: Path, values: dict[str,str]):
    """Replace selected dotenv values atomically while preserving unrelated lines."""
    original = path.read_text(encoding='utf-8') if path.exists() else ''
    lines = original.splitlines()
    written = set()
    output = []
    for line in lines:
        match = re.match(r'^(\s*)(?:export\s+)?([A-Za-z_][A-Za-z0-9_]*)\s*=.*$', line)
        if match and match.group(2) in values:
            key = match.group(2)
            if key not in written:
                output.append(f'{key}={values[key]}')
                written.add(key)
        else:
            output.append(line)
    remaining = {key:value for key,value in values.items() if key not in written}
    if output and output[-1] != '' and remaining:
        output.append('')
    output.extend(f'{key}={value}' for key,value in remaining.items())
    content = '\n'.join(output) + ('\n' if output else '')
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name+'.', suffix='.tmp', dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8', newline='\n') as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)

@app.post('/api/settings')
def update_settings(body: SettingsUpdate, request: Request):
    if request.state.identity['role'] != 'admin':
        fail('仅管理员可以修改模型配置。',403)
    base_url = body.base_url.strip().rstrip('/')
    parsed = urlparse(base_url)
    if parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password:
        fail('服务地址必须是有效的 HTTPS 地址。')
    if parsed.query or parsed.fragment:
        fail('服务地址不能包含查询参数或片段。')
    model = body.model.strip()
    embedding_model = body.embedding_model.strip()
    if not _MODEL_NAME.fullmatch(model) or not _MODEL_NAME.fullmatch(embedding_model):
        fail('模型名称格式无效。')
    api_key = body.api_key.strip()
    if api_key and not _API_KEY.fullmatch(api_key):
        fail('API Key 格式或长度无效。')
    values = {
        'RADAR_BASE_URL': base_url,
        'RADAR_MODEL': model,
        'RADAR_EMBEDDING_MODEL': embedding_model,
        'RADAR_PAID_ENABLED': str(body.paid_enabled).lower(),
    }
    if api_key:
        values['RADAR_API_KEY'] = api_key
    try:
        _replace_env_values(_env_path(), values)
    except OSError:
        fail('配置文件保存失败，请检查文件权限。', 500)
    os.environ.update(values)
    return _settings_detail()

def _visible_project(connection, project_id, identity, write=False):
    row = connection.execute('SELECT * FROM projects WHERE id=?', (project_id,)).fetchone()
    visible = row and (row['owner_type']=='system' or (row['owner_type']=='user' and row['owner_user_id']==identity['id']))
    if not visible:
        fail('项目不存在。',404)
    if write and row['owner_type']=='system':
        fail('项目不存在。',404)
    return row

def visible_project_or_404(project_id, identity, write=False):
    with db.connect() as connection:
        return dict(_visible_project(connection, project_id, identity, write))

def _project_public(row, identity):
    return {
        'id':row['id'],
        'name':row['name'],
        'description':row['description'],
        'created_at':row['created_at'],
        'owner_type':row['owner_type'],
        'is_owner':row['owner_type']=='user' and row['owner_user_id']==identity['id'],
        'source_project_id':row['source_project_id'],
    }

class Project(BaseModel):
    name: str = Field(min_length=1,max_length=100)
    description: str = Field(default='',max_length=2000)

@app.get('/api/projects')
def projects(request: Request):
    identity = request.state.identity
    with db.connect() as c:
        rows = c.execute("SELECT * FROM projects WHERE owner_type='system' OR (owner_type='user' AND owner_user_id=?) ORDER BY created_at DESC", (identity['id'],))
        return [_project_public(row, identity) for row in rows]

@app.post('/api/projects')
def create_project(body: Project, request: Request):
    if not body.name.strip(): fail('项目名称不能为空。')
    identity = request.state.identity
    record=dict(id=db.uid(),name=body.name.strip(),description=body.description,created_at=time.time(),
                owner_type='user',owner_user_id=identity['id'],source_project_id=None)
    with db.connect() as c:
        c.execute('''INSERT INTO projects(id,name,description,created_at,owner_type,owner_user_id,source_project_id)
            VALUES (:id,:name,:description,:created_at,:owner_type,:owner_user_id,:source_project_id)''',record)
    return _project_public(record, identity)

@app.post('/api/projects/{pid}/copy')
def copy_project(pid: str, request: Request):
    identity = request.state.identity
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        source = _visible_project(c, pid, identity)
        if source['owner_type'] != 'system':
            fail('仅系统样本可以复制。')
        existing = c.execute("SELECT * FROM projects WHERE owner_type='user' AND owner_user_id=? AND source_project_id=? ORDER BY created_at LIMIT 1",
                             (identity['id'],pid)).fetchone()
        if existing:
            return _project_public(existing, identity)
        new_id = db.uid()
        now = time.time()
        c.execute('''INSERT INTO projects(id,name,description,created_at,owner_type,owner_user_id,source_project_id)
            VALUES (?,?,?,?,?,?,?)''',
            (new_id,source['name']+'（我的副本）',source['description'],now,'user',identity['id'],source['id']))
        batch_ids = {}
        for batch in c.execute('SELECT * FROM batches WHERE project_id=? ORDER BY created_at', (pid,)).fetchall():
            copied_batch_id = db.uid()
            batch_ids[batch['id']] = copied_batch_id
            c.execute('''INSERT INTO batches(id,project_id,name,row_count,quality,created_at,data_kind)
                VALUES (?,?,?,?,?,?,?)''',
                (copied_batch_id,new_id,batch['name'],batch['row_count'],batch['quality'],batch['created_at'],batch['data_kind']))
        for review in c.execute('SELECT * FROM reviews WHERE project_id=?', (pid,)).fetchall():
            review_id = db.uid()
            payload = json.loads(review['payload'])
            payload['id'] = review_id
            c.execute('''INSERT INTO reviews(id,project_id,batch_id,fingerprint,payload) VALUES (?,?,?,?,?)''',
                (review_id,new_id,batch_ids.get(review['batch_id']),review['fingerprint'],db.dump(payload)))
        copied = c.execute('SELECT * FROM projects WHERE id=?', (new_id,)).fetchone()
        return _project_public(copied, identity)

def parse_file(raw, filename, sheet=''):
    if len(raw)>10*1024*1024: fail('文件不能超过10MB。')
    sheets=[]
    if filename.lower().endswith('.xlsx'):
        import zipfile
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>60*1024*1024: fail('Excel解压后过大。')
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            sheets=workbook.sheetnames
            if sheet and sheet not in sheets: fail('工作表不存在。')
            ws=workbook[sheet or sheets[0]]
            values=ws.iter_rows(values_only=True)
            columns=[str(x or '').strip() for x in next(values,[])]
            rows=[]
            for cells in values:
                if any(v is not None for v in cells): rows.append({k:str(v) if v is not None else '' for k,v in zip(columns,cells)})
                if len(rows)>MAX_ROWS: fail('最多导入5000条。')
            workbook.close()
        except HTTPException: raise
        except Exception: fail('Excel无法读取，请使用有效的.xlsx文件。')
    elif filename.lower().endswith('.csv'):
        try: text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: text=raw.decode('gb18030')
            except UnicodeDecodeError: fail('CSV编码无法识别。')
        reader=csv.DictReader(io.StringIO(text)); columns=reader.fieldnames or []; rows=[]
        for row in reader:
            if None in row: fail('CSV行列数量不一致。')
            rows.append({k:v or '' for k,v in row.items()})
            if len(rows)>MAX_ROWS: fail('最多导入5000条。')
    else: fail('仅支持.csv与.xlsx。')
    if not columns or any(not x for x in columns) or len(columns)!=len(set(columns)): fail('列名不能为空或重复。')
    if len(columns)>100: fail('最多100列。')
    return columns,rows,sheets

@app.post('/api/projects/{pid}/imports/preview')
async def preview(pid: str, request: Request, file: UploadFile=File(...), sheet: str=Form('')):
    with db.connect() as c: _visible_project(c,pid,request.state.identity,write=True)
    raw=await file.read(10*1024*1024+1)
    columns,rows,sheets=parse_file(raw,file.filename or '',sheet)
    preview_id=db.uid()
    with db.connect() as c:
        c.execute('DELETE FROM previews WHERE created_at<?',(time.time()-86400,))
        c.execute('INSERT INTO previews VALUES (?,?,?,?)',(preview_id,pid,db.dump({'columns':columns,'rows':rows}),time.time()))
    return dict(preview_id=preview_id,columns=columns,sheets=sheets,rows=rows[:10],total_rows=len(rows))

class Import(BaseModel):
    preview_id: str
    mapping: dict[str,str]
    source_name: str = Field(default='导入数据',max_length=200)
    data_kind: str = Field(default='用户导入，来源未验证',max_length=200)

def import_rows(pid, rows, mapping, name, kind, identity=None):
    if not mapping.get('text'): fail('必须映射评论列。')
    if len(set(mapping.values()))!=len(mapping): fail('同一列不能映射多个字段。')
    allowed={'text','brand','product','date','platform','source_url','id'}
    if set(mapping)-allowed: fail('存在不支持的字段。')
    batch_id=db.uid(); kept=0; blanks=0; duplicates=0
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        if identity is None:
            if not c.execute('SELECT 1 FROM projects WHERE id=?',(pid,)).fetchone():
                fail('项目不存在。',404)
        else:
            _visible_project(c,pid,identity,write=True)
        c.execute('INSERT INTO batches VALUES (?,?,?,?,?,?,?)',(batch_id,pid,name,0,'{}',time.time(),kind))
        for source in rows:
            row={key:str(source.get(col,'')).strip() for key,col in mapping.items()}
            if not row.get('text'): blanks+=1; continue
            if len(row['text'])>4000: fail('评论不能超过4000字。')
            for key,value in {'brand':'未指明','product':'未指明','date':'','platform':'未提供','source_url':''}.items(): row[key]=row.get(key) or value
            if row['source_url'] and urlparse(row['source_url']).scheme not in ('http','https'): row['source_url']=''
            fingerprint=hashlib.sha256(db.dump([re.sub(r'\s+','',row['text']),row['brand'],row['product'],row['platform'],row['date']]).encode()).hexdigest()
            row['source_id']=row.pop('id',''); row['id']=db.uid(); row['data_kind']=kind
            cur=c.execute('INSERT OR IGNORE INTO reviews VALUES (?,?,?,?,?)',(row['id'],pid,batch_id,fingerprint,db.dump(row)))
            kept+=cur.rowcount; duplicates+=1-cur.rowcount
        quality=dict(input=len(rows),kept=kept,removed=blanks+duplicates,blank=blanks,duplicates=duplicates)
        c.execute('UPDATE batches SET row_count=?,quality=? WHERE id=?',(kept,db.dump(quality),batch_id))
        record=dict(c.execute('SELECT * FROM batches WHERE id=?',(batch_id,)).fetchone()); record['quality']=quality
    return record

@app.post('/api/projects/{pid}/imports')
def commit_import(pid: str, body: Import, request: Request):
    with db.connect() as c:
        _visible_project(c,pid,request.state.identity,write=True)
        row=c.execute('SELECT payload FROM previews WHERE id=? AND project_id=? AND created_at>?',(body.preview_id,pid,time.time()-86400)).fetchone()
    if not row: fail('预览不存在或已过期。',404)
    payload=json.loads(row[0])
    if set(body.mapping.values())-set(payload['columns']): fail('映射列不存在。')
    return import_rows(pid,payload['rows'],body.mapping,body.source_name,body.data_kind,request.state.identity)

@app.get('/api/projects/{pid}/imports')
def imports(pid: str, request: Request):
    with db.connect() as c:
        _visible_project(c,pid,request.state.identity); rows=[dict(r) for r in c.execute('SELECT * FROM batches WHERE project_id=? ORDER BY created_at DESC',(pid,))]
    for row in rows: row['quality']=json.loads(row['quality'])
    return rows

@app.post('/api/projects/{pid}/demo')
def demo(pid: str, request: Request):
    from pathlib import Path
    path=Path(__file__).resolve().parent.parent/'data/demo_reviews.csv'
    columns,rows,_=parse_file(path.read_bytes(),path.name)
    return import_rows(pid,rows,{k:k for k in ('text','brand','product','date','platform','source_url','id') if k in columns},'合成演示样本','合成数据 · 非真实消费者反馈',request.state.identity)

@app.get('/api/projects/{pid}/reviews')
def reviews(pid: str,request: Request,batch_id: str='',q: str='',offset: int=0,limit: int=50):
    if offset<0 or limit<1 or limit>200: fail('分页参数无效。')
    with db.connect() as c:
        _visible_project(c,pid,request.state.identity)
        rows=[json.loads(r[0]) for r in c.execute('SELECT payload FROM reviews WHERE project_id=? AND (?="" OR batch_id=?)',(pid,batch_id,batch_id))]
    if q:
        needle=q.casefold()
        rows=[r for r in rows if needle in r['text'].casefold() or needle in str(r.get('source_id','')).casefold()]
    return {'items':rows[offset:offset+limit],'total':len(rows)}

class Job(BaseModel):
    batch_ids: list[str]
    mode: str='offline'
    clusters: int=Field(default=8,ge=1,le=20)
    consent: bool=False

@app.post('/api/projects/{pid}/jobs')
def create_job(pid: str,body: Job,request: Request):
    if body.mode not in ('offline','llm'): fail('分析模式无效。')
    if body.mode=='llm' and not (body.consent and _settings_detail()['configured'] and _settings_detail()['paid_enabled']): fail('真实模型需完整配置、服务端启用与本次同意。')
    if not body.batch_ids: fail('请选择导入批次。')
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE'); _visible_project(c,pid,request.state.identity,write=True)
        batches=list(c.execute('SELECT id FROM batches WHERE project_id=?',(pid,)))
        if set(body.batch_ids)-{b['id'] for b in batches}: fail('批次不属于本项目。')
        selected=[r['id'] for r in c.execute('SELECT id,batch_id FROM reviews WHERE project_id=?',(pid,)) if r['batch_id'] in body.batch_ids]
        if not selected or len(selected)>MAX_ROWS: fail('任务需包含1至5000条评论。')
        params=db.dump({'review_ids':selected,'clusters':body.clusters})
        previous=c.execute("SELECT * FROM jobs WHERE project_id=? AND mode=? AND parameters=? AND status IN ('queued','running')",(pid,body.mode,params)).fetchone()
        if previous: return db.job_public(previous)
        jid=db.uid(); now=time.time()
        c.execute('INSERT INTO jobs(id,project_id,status,stage,total,mode,created_at,updated_at,parameters) VALUES (?,?,?,?,?,?,?,?,?)',(jid,pid,'queued','等待处理',len(selected),body.mode,now,now,params))
        return db.job_public(c.execute('SELECT * FROM jobs WHERE id=?',(jid,)).fetchone())

def get_job(c,jid,identity,write=False):
    row=c.execute('SELECT * FROM jobs WHERE id=?',(jid,)).fetchone()
    if not row: fail('任务不存在。',404)
    _visible_project(c,row['project_id'],identity,write=write)
    return row

@app.get('/api/projects/{pid}/jobs')
def jobs(pid: str, request: Request):
    with db.connect() as c:
        _visible_project(c,pid,request.state.identity); return [db.job_public(r) for r in c.execute('SELECT * FROM jobs WHERE project_id=? ORDER BY created_at DESC',(pid,))]

@app.get('/api/jobs/{jid}')
def job(jid: str, request: Request):
    with db.connect() as c: return db.job_public(get_job(c,jid,request.state.identity))

@app.post('/api/jobs/{jid}/cancel')
def cancel(jid: str, request: Request):
    with db.connect() as c:
        row=get_job(c,jid,request.state.identity,write=True)
        if row['status'] in ('queued','running'):
            c.execute("UPDATE jobs SET cancel=1,status='cancelled',stage='已取消',updated_at=? WHERE id=?",(time.time(),jid))
        return db.job_public(get_job(c,jid,request.state.identity,write=True))

@app.post('/api/jobs/{jid}/retry')
def retry(jid: str, request: Request):
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE'); row=get_job(c,jid,request.state.identity,write=True)
        if row['status'] not in ('failed','cancelled'): fail('仅失败或取消的任务可重试。')
        if row['mode']=='llm' and not (_settings_detail()['configured'] and _settings_detail()['paid_enabled']): fail('真实模型尚未启用。')
        c.execute("UPDATE jobs SET status='queued',stage='等待续跑',error='',cancel=0,lease=NULL,lease_until=0,updated_at=? WHERE id=?",(time.time(),jid))
        return db.job_public(get_job(c,jid,request.state.identity,write=True))

def result_for(c,jid,identity,write=False):
    row=get_job(c,jid,identity,write=write)
    if row['status']!='succeeded' or not row['result']: fail('分析结果尚未完成。',409)
    return row,json.loads(row['result'])

@app.get('/api/jobs/{jid}/result')
def result(jid: str, request: Request):
    with db.connect() as c: return result_for(c,jid,request.state.identity)[1]

class Card(BaseModel):
    hypothesis: str=Field(min_length=1,max_length=4000)
    reviewed: bool

@app.patch('/api/jobs/{jid}/cards/{cid}')
def edit_card(jid: str,cid: str,body: Card,request: Request):
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE'); _,data=result_for(c,jid,request.state.identity,write=True)
        card=next((x for x in data['cards'] if x['id']==cid),None)
        if not card: fail('机会卡不存在。',404)
        try:
            from .engine import normalized_text
        except ImportError:
            from engine import normalized_text
        texts={normalized_text(e['text']) for e in card['evidence']}
        if body.reviewed and len(texts)<3: fail('证据不足3条独立原文。')
        card.update(hypothesis=body.hypothesis,reviewed=body.reviewed,reviewed_at=time.time() if body.reviewed else None)
        c.execute('UPDATE jobs SET result=? WHERE id=?',(db.dump(data),jid))
    return card

class Report(BaseModel):
    title: str=Field(default='消费者洞察报告',min_length=1,max_length=200)
    draft: bool=True

@app.post('/api/jobs/{jid}/reports')
def create_report(jid: str,body: Report,request: Request):
    with db.connect() as c:
        row,data=result_for(c,jid,request.state.identity,write=True)
        cards=data['cards'] if body.draft else [x for x in data['cards'] if x.get('reviewed')]
        if not body.draft and not cards: fail('正式报告至少需要一张已复核机会卡。')
        snapshot={**data,'cards':cards}; rid=db.uid(); now=time.time()
        provenance=data.get('provenance',{})
        kinds=provenance.get('data_kinds',{})
        kind_text='；'.join(f'{kind}：{count} 条' for kind,count in sorted(kinds.items())) or '未记录'
        quality=data.get('quality',{})
        lines=['# '+body.title,'','## 变更定位','本报告为不可变分析快照。','',
               '## 1. 报告身份与边界',
               '状态：'+('草稿，机会卡尚待人工复核' if body.draft else '所含机会卡已人工复核'),
               '数据类型：'+kind_text,
               '来源状态：'+provenance.get('source_status','数据来源未独立核验。'),
               f"样本数量：{quality.get('input_count',0)} 条；规范化去重后：{quality.get('unique_evidence_count',0)} 条。",
               '结论仅针对导入样本；机会为待验证假设，不代表全市场、总体消费者或产品因果。','',
               '## 2. 分析方法与限制','方法：'+data.get('method','未记录'),
               '趋势含义：若日期满足比较条件，仅表示该导入样本内固定主题份额变化，不能解释为市场增长。',
               '分类边界：固定主题用于稳定统计；向量聚类（如启用）仅供探索，不能自动生成经证实的业务结论。','']
        for card in cards:
            lines+=['## 3. 机会卡：'+card['topic'],
                    '**样本事实**：'+card.get('observed_fact','未记录'),
                    '**待验证假设**：'+card['hypothesis'],
                    '**验证建议**：'+card.get('validation_plan','先复核证据并补充用户研究。'),
                    '**结论边界**：'+card.get('boundary','仅代表导入样本。'),'','### 原始证据']
            for e in card['evidence']:
                evidence_meta='；'.join(filter(None,[
                    '证据编号：'+e['id'],
                    ('平台：'+e.get('platform','')) if e.get('platform') else '',
                    ('日期：'+e.get('date','')) if e.get('date') else '',
                    ('产品：'+e.get('product','')) if e.get('product') else '',
                ]))
                lines+=['> '+e['text'].replace('\n','\n> '),evidence_meta,'']
        markdown='\n'.join(lines)
        c.execute('INSERT INTO reports VALUES (?,?,?,?,?,?,?,?)',(rid,row['project_id'],jid,body.title,now,int(body.draft),db.dump(snapshot),markdown))
    return {'id':rid,'title':body.title,'created_at':now,'draft':body.draft,'job_id':jid,'markdown':markdown}

@app.get('/api/projects/{pid}/reports')
def reports(pid: str, request: Request):
    with db.connect() as c:
        _visible_project(c,pid,request.state.identity); return [{**dict(r),'draft':bool(r['draft'])} for r in c.execute('SELECT id,title,created_at,draft,job_id FROM reports WHERE project_id=? ORDER BY created_at DESC',(pid,))]

def _report_for(c, rid, identity):
    row=c.execute('SELECT * FROM reports WHERE id=?',(rid,)).fetchone()
    if not row: fail('报告不存在。',404)
    _visible_project(c,row['project_id'],identity)
    return row

@app.get('/api/reports/{rid}')
def report(rid: str, request: Request):
    with db.connect() as c: row=_report_for(c,rid,request.state.identity)
    return {**dict(row),'draft':bool(row['draft']),'payload':json.loads(row['payload'])}

@app.get('/api/reports/{rid}/download')
def download_report(rid: str, request: Request):
    with db.connect() as c: row=_report_for(c,rid,request.state.identity)
    return Response(row['markdown'],media_type='text/markdown; charset=utf-8',headers={'Content-Disposition':'attachment; filename="report.md"'})

def csv_response(rows,fields,filename):
    output=io.StringIO(); writer=csv.DictWriter(output,fieldnames=fields,extrasaction='ignore'); writer.writeheader()
    for row in rows:
        clean={k:json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else str(v) if v is not None else '' for k,v in row.items()}
        writer.writerow({k:"'"+v if v.lstrip().startswith(('=','+','-','@','\t','\r')) else v for k,v in clean.items()})
    return Response(('\ufeff'+output.getvalue()).encode(),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="{filename}"'})

@app.get('/api/jobs/{jid}/export.csv')
def export(jid: str, request: Request):
    with db.connect() as c: data=result_for(c,jid,request.state.identity)[1]
    rows=data['reviews']; return csv_response(rows,list(rows[0]) if rows else ['id','text'],'reviews.csv')

@app.get('/api/jobs/{jid}/annotation.csv')
def annotation(jid: str, request: Request):
    import random
    try:
        from .engine import normalized_text
    except ImportError:
        from engine import normalized_text
    unique={}
    with db.connect() as c: data=result_for(c,jid,request.state.identity)[1]
    for row in data['reviews']:
        unique.setdefault(normalized_text(row['text']),row)
    candidates=list(unique.values())
    sampled=random.Random(42).sample(candidates,min(100,len(candidates)))
    rows=[{'id':r['id'],'text':r['text']} for r in sampled]
    return csv_response(rows,['id','text','feedback_type','dimension','sentiment','severity'],'annotation.csv')

@app.post('/api/jobs/{jid}/evaluate')
async def evaluate(jid: str,request: Request,file: UploadFile=File(...)):
    try:
        from .engine import evaluate_reviews
    except ImportError: from engine import evaluate_reviews
    _,gold,_=parse_file(await file.read(10*1024*1024+1),file.filename or '')
    with db.connect() as c: data=result_for(c,jid,request.state.identity)[1]
    try: return evaluate_reviews(data['reviews'],gold)
    except ValueError as exc: fail(str(exc))


@app.get('/{path:path}', include_in_schema=False)
def frontend(path: str):
    if path == 'api' or path.startswith('api/'):
        fail('接口不存在。',404)
    root = Path(os.getenv('RADAR_FRONTEND_DIR', str(Path(__file__).resolve().parent.parent / 'frontend/dist'))).resolve()
    target = (root / path).resolve()
    if not target.is_relative_to(root):
        fail('路径不可访问。',404)
    if target.is_file():
        return FileResponse(target)
    if Path(path).suffix:
        fail('资源不存在。',404)
    if (root / 'index.html').is_file():
        return FileResponse(root / 'index.html')
    fail('前端尚未构建，请先执行 npm run build。',503)
